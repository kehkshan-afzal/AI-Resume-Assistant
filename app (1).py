import io
import json
import os

import streamlit as st
from google import genai
from google.genai import types
from docx import Document

APP_TITLE = "Resume ATS Analyzer"
MODEL_NAME = "gemini-2.5-flash"


def get_api_key():
    """Read the Gemini API key from Streamlit secrets or an environment variable."""
    try:
        key = st.secrets.get("GEMINI_API_KEY")
    except Exception:
        key = None
    return key or os.getenv("GEMINI_API_KEY")


def extract_docx_text(file_bytes):
    """Extract readable text from a DOCX resume."""
    document = Document(io.BytesIO(file_bytes))
    parts = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            row_text = " | ".join(cell for cell in cells if cell)
            if row_text:
                parts.append(row_text)

    return "\n".join(parts)


def analyze_resume(uploaded_file, api_key):
    """Send the resume to Gemini and return structured ATS analysis."""
    client = genai.Client(api_key=api_key)
    file_bytes = uploaded_file.getvalue()
    mime_type = uploaded_file.type or "application/octet-stream"

    prompt = """
You are an expert resume parser and ATS-optimization reviewer.

Analyze the uploaded resume for general ATS compatibility. Do NOT pretend this is
the score produced by a specific commercial ATS. Instead, provide an estimated
ATS-readiness score from 0 to 100 based only on the resume itself.

Evaluate these areas:
1. Contact information and header clarity
2. Standard section headings
3. Formatting and ATS readability
4. Work experience clarity and bullet quality
5. Quantifiable achievements
6. Skills section quality
7. Education/certification clarity
8. Keyword usage and relevance
9. Consistency, grammar, and professionalism
10. Potential ATS parsing problems such as columns, tables, graphics, icons,
   unusual symbols, headers/footers, or overly complex layouts

Important:
- Never invent experience, skills, education, employers, dates, or achievements.
- Identify missing information only as a recommendation.
- Give practical improvements that the applicant can actually make.
- If the resume appears to contain a design element that may hurt parsing,
  explain why.
- Keep the score explainable: each category should have a numeric score.
- Return valid JSON matching the requested schema.
"""

    category_names = [
        "contact_and_header",
        "section_structure",
        "ats_readability",
        "experience",
        "quantifiable_achievements",
        "skills",
        "education_and_certifications",
        "keywords",
        "consistency_and_grammar",
        "formatting_risks",
    ]

    schema = {
        "type": "OBJECT",
        "properties": {
            "ats_score": {"type": "INTEGER"},
            "summary": {"type": "STRING"},
            "category_scores": {
                "type": "OBJECT",
                "properties": {
                    name: {"type": "INTEGER"} for name in category_names
                },
                "required": category_names,
            },
            "strengths": {
                "type": "ARRAY",
                "items": {"type": "STRING"},
            },
            "improvements": {
                "type": "ARRAY",
                "items": {"type": "STRING"},
            },
            "ats_risks": {
                "type": "ARRAY",
                "items": {"type": "STRING"},
            },
            "priority_actions": {
                "type": "ARRAY",
                "items": {"type": "STRING"},
            },
        },
        "required": [
            "ats_score",
            "summary",
            "category_scores",
            "strengths",
            "improvements",
            "ats_risks",
            "priority_actions",
        ],
    }

    config = types.GenerateContentConfig(
        temperature=0.2,
        response_mime_type="application/json",
        response_schema=schema,
    )

    if mime_type == "application/pdf":
        uploaded = client.files.upload(
            file=io.BytesIO(file_bytes),
            config={"mime_type": "application/pdf"},
        )
        response = client.models.generate_content(
            model=MODEL_NAME,
            contents=[prompt, uploaded],
            config=config,
        )
    elif (
        mime_type
        == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    ):
        resume_text = extract_docx_text(file_bytes)
        if not resume_text.strip():
            raise ValueError("The DOCX file does not contain readable text.")
        response = client.models.generate_content(
            model=MODEL_NAME,
            contents=[prompt, "\nRESUME TEXT:\n" + resume_text],
            config=config,
        )
    else:
        raise ValueError("Unsupported file type. Please upload a PDF or DOCX.")

    if not response.text:
        raise ValueError("Gemini returned an empty response.")

    return json.loads(response.text)


def score_label(score):
    if score >= 85:
        return "Strong ATS readiness"
    if score >= 70:
        return "Good ATS readiness"
    if score >= 50:
        return "Needs improvement"
    return "Major improvements recommended"


st.set_page_config(
    page_title=APP_TITLE,
    page_icon="📄",
    layout="wide",
)

st.title("📄 Resume ATS Analyzer")
st.write(
    "Upload your resume to get an estimated ATS-readiness score, "
    "ATS risks, strengths, and practical improvement suggestions."
)

st.info(
    "This is an AI-based estimate, not an official score from a specific ATS. "
    "Different applicant-tracking systems use different parsing and ranking rules."
)

uploaded_file = st.file_uploader(
    "Upload your resume",
    type=["pdf", "docx"],
    help="PDF or DOCX only. Avoid uploading sensitive information you do not want processed by an external AI service.",
)

if uploaded_file:
    st.caption(f"Selected: {uploaded_file.name}")

    if st.button("🔍 Analyze Resume", type="primary", use_container_width=True):
        api_key = get_api_key()

        if not api_key:
            st.error(
                "Gemini API key not found. Add GEMINI_API_KEY to Streamlit Secrets "
                "or set it as an environment variable."
            )
            st.stop()

        with st.spinner("Analyzing your resume with Gemini Flash..."):
            try:
                result = analyze_resume(uploaded_file, api_key)

                score = max(0, min(100, int(result["ats_score"])))

                st.divider()
                st.subheader("ATS-readiness score")

                col1, col2 = st.columns([1, 2])
                with col1:
                    st.metric("Estimated Score", f"{score}/100")
                with col2:
                    st.progress(score / 100)
                    st.write(score_label(score))

                st.subheader("Summary")
                st.write(result["summary"])

                st.subheader("Category scores")
                scores = result["category_scores"]

                display_names = {
                    "contact_and_header": "Contact & Header",
                    "section_structure": "Section Structure",
                    "ats_readability": "ATS Readability",
                    "experience": "Experience",
                    "quantifiable_achievements": "Quantifiable Achievements",
                    "skills": "Skills",
                    "education_and_certifications": "Education & Certifications",
                    "keywords": "Keywords",
                    "consistency_and_grammar": "Consistency & Grammar",
                    "formatting_risks": "Formatting Risks",
                }

                score_cols = st.columns(2)
                for index, (key, label) in enumerate(display_names.items()):
                    with score_cols[index % 2]:
                        value = max(0, min(100, int(scores.get(key, 0))))
                        st.write(f"**{label}: {value}/100**")
                        st.progress(value / 100)

                left, right = st.columns(2)

                with left:
                    st.subheader("✅ Strengths")
                    for item in result["strengths"]:
                        st.write(f"• {item}")

                    st.subheader("⚠️ ATS risks")
                    for item in result["ats_risks"]:
                        st.write(f"• {item}")

                with right:
                    st.subheader("🛠️ Improvements")
                    for item in result["improvements"]:
                        st.write(f"• {item}")

                    st.subheader("🎯 Priority actions")
                    for index, item in enumerate(result["priority_actions"], start=1):
                        st.write(f"{index}. {item}")

            except json.JSONDecodeError:
                st.error("Gemini returned an invalid analysis format. Please try again.")
            except Exception as exc:
                st.error(f"Analysis failed: {exc}")
else:
    st.markdown(
        """
        ### How it works
        1. Upload a PDF or DOCX resume.
        2. Gemini Flash reviews the resume.
        3. The app calculates an estimated ATS-readiness score.
        4. You receive category scores, ATS risks, strengths, and improvements.
        """
    )
